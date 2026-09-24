#!/usr/bin/env python3
"""生成示例课件 示例课件.pptx（5 页，16:9），供 slides-to-handout 的完整示例使用。

运行: python3 make_demo_deck.py   # 需要 python-pptx
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

DARK = RGBColor(0x1A, 0x1A, 0x2E)
BLUE = RGBColor(0x4A, 0x4A, 0x8A)
GRAY = RGBColor(0x66, 0x66, 0x66)
FONT = "PingFang SC"

SLIDES = [
    {
        "kind": "cover",
        "title": "股票估值入门",
        "subtitle": "从股利贴现模型到戈登增长模型",
        "note": "示例课件 · 共 5 页 · 用于演示 slides-to-handout",
    },
    {
        "title": "股利贴现模型（DDM）",
        "bullets": [
            "股票的价值 = 未来全部股利的现值之和",
            "假设：投资者长期持有，股利是股东能拿到的唯一现金流",
            "P0 = Σ(t=1→∞) Div_t / (1 + r_E)^t",
            "难点：无限期股利无法直接预测，需要加假设",
        ],
    },
    {
        "title": "戈登增长模型（Gordon Growth Model）",
        "bullets": [
            "假设股利以固定增长率 g 永续增长：Div_t = Div_1 × (1 + g)^(t-1)",
            "代入 DDM，得到等比级数",
            "P0 = Div_1 / (r_E − g)",
            "适用条件：r_E > g，且 g 长期可持续",
        ],
    },
    {
        "title": "例题：用戈登模型估值",
        "bullets": [
            "某公司明年每股股利 Div_1 = 2.00 元",
            "股权资本成本 r_E = 10%，股利永续增长率 g = 4%",
            "P0 = 2.00 / (0.10 − 0.04) = 33.33 元",
            "若当前股价 30 元，是否低估？",
        ],
    },
    {
        "title": "模型局限",
        "bullets": [
            "当 g → r_E 时价格趋于无穷，估值对 g 极其敏感",
            "只适合分红稳定、增长平稳的公司",
            "高增长公司需要两阶段或多阶段模型",
            "长期看 g 不可能持续高于名义 GDP 增速",
        ],
    },
]


def add_textbox(slide, left, top, width, height, text, size, bold=False, color=DARK,
                align=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    return box


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    for item in SLIDES:
        slide = prs.slides.add_slide(blank)
        if item.get("kind") == "cover":
            add_textbox(slide, 1.0, 2.6, 11.333, 1.2, item["title"], 44, bold=True,
                        color=DARK)
            add_textbox(slide, 1.0, 3.9, 11.333, 0.8, item["subtitle"], 24, color=BLUE)
            add_textbox(slide, 1.0, 5.6, 11.333, 0.6, item["note"], 14, color=GRAY)
            continue

        add_textbox(slide, 0.9, 0.7, 11.5, 0.9, item["title"], 30, bold=True, color=BLUE)
        box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.6))
        tf = box.text_frame
        tf.word_wrap = True
        for i, b in enumerate(item["bullets"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = ("• " if not b.startswith("P0") else "") + b
            run.font.size = Pt(20 if not b.startswith("P0") else 24)
            run.font.bold = b.startswith("P0")
            run.font.color.rgb = DARK if not b.startswith("P0") else BLUE
            run.font.name = FONT
            p.space_after = Pt(14)

    prs.save("示例课件.pptx")
    print("已生成 示例课件.pptx（%d 页）" % len(SLIDES))


if __name__ == "__main__":
    main()
